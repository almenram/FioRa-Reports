from openpyxl import Workbook
from uuid import uuid4


def exportar_excel(registros, nombre_archivo=None):

    if nombre_archivo is None:
        nombre_archivo = f"reporte_{uuid4().hex[:6]}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Reportes"

    encabezados = [
        "ID",
        "Código",
        "Descripción",
        "Cantidad",
        "Solicitante",
        "Líder",
        "Fabricación",
        "Fecha"
    ]

    ws.append(encabezados)

    for registro in registros:
        ws.append([
            registro.id,
            registro.codigo,
            registro.descripcion,
            registro.cantidad,
            registro.solicitante,
            registro.lider,
            registro.fabricacion,
            registro.fecha
        ])

    for fila in range(2, ws.max_row + 1):
        ws.cell(row=fila, column=4).number_format = "0.00"

    for columna in ws.columns:
        longitud = 0

        for celda in columna:
            if celda.value is not None:
                longitud = max(longitud, len(str(celda.value)))

        ws.column_dimensions[
            columna[0].column_letter
        ].width = longitud + 2

    wb.save(nombre_archivo)

    return nombre_archivo   